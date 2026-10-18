import csv
import io
import logging
import unicodedata
import uuid
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from app.core.database import get_db
from app.models.inventario import Inventario, Producto
from app.models.marcas import Estado, Marca
from app.schemas.inventario import ImportResult, ProductoCreate, ProductoResponse, ProductoUpdate
from app.services.product_image_storage import MAX_PRODUCT_IMAGE_BYTES, upload_product_image, validate_image
from app.services.stock_service import available_stock_by_product, critical_stock_by_product

router = APIRouter()
logger = logging.getLogger("jmstore.inventory")


def _product_response(product: Producto, brand_name: str | None) -> dict:
    return {**{column.name: getattr(product, column.name) for column in Producto.__table__.columns},
            "marca_nombre": brand_name}


@router.get("/marcas")
def list_inventory_brands(db: Session = Depends(get_db)):
    return [
        {"id": brand_id, "nombre_marca": brand_name}
        for brand_id, brand_name in db.query(Marca.id, Marca.nombre_marca).order_by(Marca.nombre_marca).all()
    ]


@router.get("/estados")
def list_product_states(db: Session = Depends(get_db)):
    return [
        {"id": state_id, "nombre": name, "modulo": module}
        for state_id, name, module in db.query(Estado.id, Estado.nombre, Estado.modulo)
        .order_by(Estado.modulo, Estado.nombre).all()
    ]


@router.get("/", response_model=list[ProductoResponse])
def list_products(
    q: str = Query(default="", max_length=150),
    marca_id: int | None = None,
    db: Session = Depends(get_db),
):
    query = (db.query(Producto, Marca.nombre_marca)
             .outerjoin(Marca, Marca.id == Producto.marca_id))
    if marca_id:
        query = query.filter(Producto.marca_id == marca_id)
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.filter((Producto.nombre.ilike(term)) | (Producto.sku.ilike(term)) | (Producto.codigo_qr.ilike(term)))
    rows = query.order_by(Producto.nombre).limit(1000).all()
    product_ids = [product.id for product, _ in rows]
    stocks = available_stock_by_product(db, product_ids)
    critical = critical_stock_by_product(db, product_ids)
    return [{**_product_response(product, brand_name), "stock": stocks.get(product.id, 0),
             "stock_critico": critical.get(product.id, 5)}
            for product, brand_name in rows]


@router.get("/catalogo", response_model=list[ProductoResponse])
def list_product_catalog(
    q: str = Query(default="", max_length=150),
    marca_id: int | None = None,
    db: Session = Depends(get_db),
):
    query = db.query(Producto, Marca.nombre_marca).outerjoin(Marca, Marca.id == Producto.marca_id)
    if marca_id:
        query = query.filter(Producto.marca_id == marca_id)
    if q.strip():
        term = f"%{q.strip()}%"
        query = query.filter((Producto.nombre.ilike(term)) | (Producto.sku.ilike(term)) | (Producto.codigo_qr.ilike(term)))
    rows = query.order_by(Producto.nombre).limit(1000).all()
    product_ids = [product.id for product, _ in rows]
    stocks = available_stock_by_product(db, product_ids)
    critical = critical_stock_by_product(db, product_ids)
    return [{**_product_response(product, brand_name), "stock": stocks.get(product.id, 0),
             "stock_critico": critical.get(product.id, 5)}
            for product, brand_name in rows]


@router.post("/", response_model=ProductoResponse, status_code=status.HTTP_201_CREATED)
def create_product(payload: ProductoCreate, db: Session = Depends(get_db)):
    if not db.query(Marca).filter(Marca.id == payload.marca_id).first():
        raise HTTPException(400, "La marca seleccionada no existe")
    state_id = payload.estado_id
    if not db.query(Estado).filter(Estado.id == state_id).first():
        raise HTTPException(400, "El estado seleccionado no existe")
    sku = payload.sku.strip() if payload.sku else None
    product = Producto(
        marca_id=payload.marca_id,
        codigo_qr=(payload.codigo_qr.strip() if payload.codigo_qr else None) or sku or f"JM-{uuid.uuid4().hex[:10].upper()}",
        sku=sku,
        nombre=payload.nombre.strip(),
        descripcion=payload.descripcion.strip() if payload.descripcion else None,
        precio=payload.precio,
        costo_nio=payload.costo_nio,
        costo_envio_nio=payload.costo_envio_nio,
        margen_utilidad_pct=payload.margen_utilidad_pct,
        imagen_url=payload.imagen_url,
        estado_id=state_id,
    )
    db.add(product)
    db.flush()
    db.add(Inventario(producto_id=product.id, stock=payload.stock, stock_critico=payload.stock_critico))
    db.commit()
    db.refresh(product)
    result = _product_response(product, db.query(Marca.nombre_marca).filter(Marca.id == product.marca_id).scalar())
    result["stock"] = available_stock_by_product(db, [product.id]).get(product.id, 0)
    result["stock_critico"] = critical_stock_by_product(db, [product.id]).get(product.id, 5)
    return result


@router.put("/{product_id}", response_model=ProductoResponse)
def update_product(product_id: int, payload: ProductoUpdate, db: Session = Depends(get_db)):
    product = db.query(Producto).filter(Producto.id == product_id).first()
    if not product:
        raise HTTPException(404, "Producto no encontrado")
    changes = payload.model_dump(exclude_unset=True)
    if "marca_id" in changes and not db.query(Marca).filter(Marca.id == changes["marca_id"]).first():
        raise HTTPException(400, "La marca seleccionada no existe")
    critical = changes.pop("stock_critico", None)
    stock = changes.pop("stock", None)
    for field, value in changes.items():
        setattr(product, field, value)
    if critical is not None or stock is not None:
        inventory = db.query(Inventario).filter(Inventario.producto_id == product.id).first()
        if inventory is None:
            inventory = Inventario(producto_id=product.id, stock=0)
            db.add(inventory)
        if critical is not None:
            inventory.stock_critico = critical
        if stock is not None:
            inventory.stock = stock
    db.commit()
    db.refresh(product)
    result = _product_response(product, db.query(Marca.nombre_marca).filter(Marca.id == product.marca_id).scalar())
    result["stock"] = available_stock_by_product(db, [product.id]).get(product.id, 0)
    result["stock_critico"] = critical_stock_by_product(db, [product.id]).get(product.id, 5)
    return result


@router.post("/{product_id}/imagen", response_model=ProductoResponse)
async def upload_product_photo(product_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    product = db.query(Producto).filter(Producto.id == product_id).first()
    if not product:
        raise HTTPException(404, "Producto no encontrado")
    content = await file.read(MAX_PRODUCT_IMAGE_BYTES + 1)
    extension = validate_image(content, file.content_type)
    image_url = upload_product_image(product_id, extension, file.content_type or "", content)
    product.imagen_url = image_url
    try:
        db.commit()
    except SQLAlchemyError as exc:
        db.rollback()
        logger.exception("Could not save uploaded image URL for product_id=%s", product_id)
        raise HTTPException(500, "La imagen se subió, pero no se pudo actualizar el producto") from exc
    db.refresh(product)
    result = _product_response(product, db.query(Marca.nombre_marca).filter(Marca.id == product.marca_id).scalar())
    result["stock"] = available_stock_by_product(db, [product.id]).get(product.id, 0)
    result["stock_critico"] = critical_stock_by_product(db, [product.id]).get(product.id, 5)
    return result


@router.delete("/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_product(product_id: int, db: Session = Depends(get_db)):
    product = db.query(Producto).filter(Producto.id == product_id).first()
    if not product:
        raise HTTPException(404, "Producto no encontrado")
    db.delete(product)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "No se puede eliminar: el producto tiene movimientos o ventas asociados.") from exc


def _normalise_header(value) -> str:
    raw = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(c for c in raw if not unicodedata.combining(c)).replace(" ", "_").replace("-", "_")


HEADER_ALIASES = {
    "nombre": ("nombre", "producto", "name", "product", "descripcion_producto"),
    "sku": ("sku", "codigo", "codigo_producto", "referencia", "ref"),
    "codigo_qr": ("codigo_qr", "qr", "codigo_qr_producto"),
    "precio": ("precio", "price", "precio_venta", "pvp"),
    "costo_nio": ("costo_nio", "costo_compra", "costo", "cost"),
    "costo_envio_nio": ("costo_envio_nio", "costo_envio", "envio", "shipping"),
    "margen_utilidad_pct": ("margen_utilidad_pct", "margen_utilidad", "margen", "utilidad_pct"),
    "stock": ("stock", "existencia", "existencias", "cantidad", "inventario"),
    "stock_critico": ("stock_critico", "stock_minimo", "minimo", "alerta_stock"),
    "descripcion": ("descripcion", "detalle", "observaciones"),
    "marca": ("marca", "brand", "marca_id"),
    "imagen_url": ("imagen_url", "imagen", "foto", "image_url"),
    "estado_id": ("estado_id", "estado"),
}


def _read_rows(filename: str, content: bytes) -> list[tuple[int, list]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        return [(i, list(row)) for i, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
    if filename.lower().endswith(".csv"):
        decoded = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [(i, row) for i, row in enumerate(csv.reader(io.StringIO(decoded), dialect), start=1)]
    raise HTTPException(400, "Usa un archivo CSV o Excel (.xlsx)")


@router.post("/importar", response_model=ImportResult)
async def import_products(
    file: UploadFile = File(...),
    marca_id: int | None = Form(default=None),
    db: Session = Depends(get_db),
):
    raise HTTPException(405, "Registra los productos desde una compra a un proveedor")
    content = await file.read(10 * 1024 * 1024 + 1)
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "El archivo no puede superar 10 MB")
    try:
        rows = _read_rows(file.filename or "", content)
    except (UnicodeDecodeError, ValueError) as exc:
        raise HTTPException(400, f"No se pudo leer el archivo: {exc}") from exc
    if len(rows) < 2:
        raise HTTPException(400, "El archivo debe incluir encabezados y al menos una fila")

    headers = [_normalise_header(value) for value in rows[0][1]]
    fields = {field: next((headers.index(alias) for alias in aliases if alias in headers), None)
              for field, aliases in HEADER_ALIASES.items()}
    if fields["nombre"] is None or fields["precio"] is None:
        raise HTTPException(400, "Incluye columnas nombre y precio. Se aceptan también producto y price.")
    if fields["marca"] is None and marca_id is None:
        raise HTTPException(400, "Incluye una columna marca o selecciona una marca para todo el archivo")

    default_state = db.query(Estado).filter(Estado.id == 1).first() or db.query(Estado).order_by(Estado.id).first()
    if not default_state:
        default_state = Estado(modulo="Inventario", nombre="Activo", descripcion="Estado inicial de productos")
        db.add(default_state)
        db.flush()
    created = updated = 0
    errors: list[str] = []
    for line_number, row in rows[1:]:
        if not any(value is not None and str(value).strip() for value in row):
            continue
        try:
            def value(field, fallback=None):
                index = fields.get(field)
                return row[index] if index is not None and index < len(row) else fallback

            name = str(value("nombre", "") or "").strip()
            sku_raw = value("sku")
            sku = str(sku_raw).strip() if sku_raw is not None and str(sku_raw).strip() else None
            qr_raw = value("codigo_qr")
            codigo_qr = str(qr_raw).strip() if qr_raw is not None and str(qr_raw).strip() else sku
            price_raw = value("precio")
            quantity = int(float(value("stock", 0) or 0))
            minimum = int(float(value("stock_critico", 5) or 0))
            brand_raw = value("marca", marca_id)
            if not name:
                raise ValueError("falta nombre")
            if quantity < 0 or minimum < 0:
                raise ValueError("stock no puede ser negativo")
            if brand_raw in (None, ""):
                raise ValueError("falta marca")
            if isinstance(brand_raw, (int, float)) or str(brand_raw).strip().isdigit():
                brand = db.query(Marca).filter(Marca.id == int(brand_raw)).first()
            else:
                brand = db.query(Marca).filter(Marca.nombre_marca.ilike(str(brand_raw).strip())).first()
            if not brand:
                raise ValueError(f"marca '{brand_raw}' no encontrada")
            product = db.query(Producto).filter(Producto.marca_id == brand.id, Producto.sku == sku).first() if sku else None
            margin_raw = value("margen_utilidad_pct", product.margen_utilidad_pct if product else 30)
            if margin_raw in (None, ""):
                margin_raw = product.margen_utilidad_pct if product else 30
            purchase_cost = Decimal(str(value("costo_nio", product.costo_nio if product else 0) or 0).replace(",", "."))
            shipping_cost = Decimal(str(value("costo_envio_nio", product.costo_envio_nio if product else 0) or 0).replace(",", "."))
            margin = Decimal(str(margin_raw).replace(",", "."))
            price = Decimal(str(price_raw).replace(",", ".")) if price_raw not in (None, "") else ((purchase_cost + shipping_cost) * (1 + margin / 100)).quantize(Decimal("0.01"))
            if not price.is_finite() or price < 0:
                raise ValueError("El precio no puede ser negativo")
            if not purchase_cost.is_finite() or purchase_cost < 0 or not shipping_cost.is_finite() or shipping_cost < 0:
                raise ValueError("Los costos no pueden ser negativos")
            if not margin.is_finite() or margin < 0:
                raise ValueError("La utilidad no puede ser negativa")
            attrs = {
                "marca_id": brand.id, "sku": sku, "nombre": name,
                "codigo_qr": codigo_qr or f"JM-{uuid.uuid4().hex[:10].upper()}",
                "precio": price, "stock": quantity, "stock_critico": minimum,
                "costo_nio": purchase_cost,
                "costo_envio_nio": shipping_cost,
                "margen_utilidad_pct": margin,
                "descripcion": str(value("descripcion") or "").strip() or None,
                "imagen_url": str(value("imagen_url") or "").strip() or None,
                "estado_id": int(value("estado_id", default_state.id) or default_state.id),
            }
            if not db.query(Estado.id).filter(Estado.id == attrs["estado_id"]).scalar():
                raise ValueError(f"estado_id '{attrs['estado_id']}' no encontrado")
            if product:
                for key, item in attrs.items():
                    setattr(product, key, item)
                updated += 1
            else:
                db.add(Producto(**attrs))
                created += 1
        except (ValueError, InvalidOperation, TypeError, OverflowError) as exc:
            errors.append(f"Fila {line_number}: {exc}")
    try:
        db.commit()
    except Exception:
        db.rollback()
        raise HTTPException(400, "No se pudieron guardar los productos. Revisa que aplicaste la actualización SQL de inventario en Supabase.")
    return {"creados": created, "actualizados": updated, "errores": errors[:100]}
